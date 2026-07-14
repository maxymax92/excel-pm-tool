"""Shared structural-QA constants and small workbook helpers."""

from __future__ import annotations

from datetime import datetime
from typing import TYPE_CHECKING

from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

if TYPE_CHECKING:
    from openpyxl.styles.dimensions import ColumnDimension
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.worksheet import Worksheet

EXPECTED_SHEETS = ["Overview", "Plan", "Items", "RAID", "Config", "Calc"]
PROTECTED_SHEETS = {"Overview", "Plan", "Calc"}
EXPECTED_TABLES = {
    "Items": ["tblItems"],
    "RAID": ["tblRAID"],
    "Config": [
        "tblStatuses",
        "tblTypes",
        "tblPriorities",
        "tblTeams",
        "tblRaidTypes",
        "tblRaidStatuses",
        "tblSeverity",
        "tblDeliveryHealth",
        "tblPeople",
    ],
}
EXPECTED_NAMES = {
    "cfgDueSoonDays",
    "cfgBlockedRedDays",
    "cfgStaleDays",
    "cfgReportDays",
    "cfgExecutiveStatusMaxLevel",
    "cfgKeyDateMaxLevel",
    "cfgComingUrgentDays",
    "cfgComingSoonDays",
    "cfgComingNearDays",
    "cfgComingHorizonDays",
    "cfgAlertSevScore",
    "cfgItemIDPrefix",
    "cfgRaidIDPrefix",
    "cfgNextItemID",
    "cfgNextRaidID",
    "lstActiveStatus",
    "lstDoneStatus",
    "lstCancelledStatus",
    "lstClosedRaid",
    "lstAlertRaid",
    "lstDecisionRaid",
    "dvStatus",
    "dvStatusActive",
    "dvStatusDone",
    "dvStatusCancelled",
    "dvTypes",
    "dvTypeLevels",
    "dvPriorities",
    "dvTeams",
    "dvRaidTypes",
    "dvRaidAlert",
    "dvRaidDecision",
    "dvRaidStatuses",
    "dvRaidClosed",
    "dvSeverity",
    "dvSeverityMinScore",
    "dvDeliveryHealth",
    "dvItemIDs",
    "dvPeople",
    "dvScopeLabels",
    "selPScope",
    "selPDepth",
    "selPScopeID",
    "selPFrom",
    "selPTo",
    "fnItemLookup",
    "fnTypeLevel",
    "fnAncestorAtLevel",
    "fnSortDate",
    "fnWbsKey",
    "fnBizDays",
    "fnHealthRAG",
    "fnDepOpen",
    "fnRefsValid",
    "fnIsDone",
    "fnIsActive",
    "fnIsCancelled",
}
ITEMS_CORE = [
    "ID",
    "Type",
    "Title",
    "Parent",
    "Priority",
    "Start",
    "Status",
    "Due",
    "Delivery Health",
    "Latest Status",
    "Owner",
]
PEOPLE_HEADERS = ["Person", "Role", "Team"]
RAID_CORE = [
    "RaidID",
    "Type",
    "Title",
    "Detail",
    "RelatedID",
    "Owner",
    "Status",
    "Prob",
    "Impact",
    "Severity",
    "Response",
    "NextReview",
]
RAID_SYSTEM = ["Score", "Scope", "Raised", "Closed", "Updated"]
ITEMS_HDR_ROW = 2
ITEMS_DATA_ROW = 3
RAID_HDR_ROW = 2
RAID_DATA_ROW = 3


def table_headers(worksheet: Worksheet, table_name: str) -> list[str]:
    """Return table headers in workbook order, independent of stray cells.

    Returns:
        The table-column names in workbook order.

    """
    return [column.name for column in worksheet.tables[table_name].tableColumns]


def validation_for(worksheet: Worksheet, coordinate: str) -> DataValidation | None:
    """Return the validation rule covering one coordinate, if present.

    Returns:
        The matching validation rule or ``None``.

    """
    return next(
        (rule for rule in worksheet.data_validations.dataValidation if coordinate in rule.sqref),
        None,
    )


def normalise_number_format(value: object) -> str:
    """Normalise spaces that Excel escapes when saving custom formats.

    Returns:
        A comparable number-format string.

    """
    return str(value).replace(r"\ ", " ")


def normalise_cached_date(got: object, want: object, epoch: datetime) -> object:
    """Convert a numeric Excel cache only when the assertion expects a date.

    Returns:
        A converted cached datetime or the original value.

    """
    if isinstance(want, datetime) and isinstance(got, (int, float)) and not isinstance(got, bool):
        try:
            converted = from_excel(got, epoch)
        except (OverflowError, TypeError, ValueError):
            return got
        if isinstance(converted, datetime):
            return converted
    return got


def formula_parentheses_balanced(formula: object) -> bool:
    """Return whether a generated formula has balanced parentheses and strings.

    Returns:
        ``True`` only for a balanced expression.

    """
    depth = 0
    in_string = False
    for character in str(formula):
        if character == '"':
            in_string = not in_string
        elif not in_string and character == "(":
            depth += 1
        elif not in_string and character == ")":
            depth -= 1
            if depth < 0:
                return False
    return depth == 0 and not in_string


def column_dimension_for(worksheet: Worksheet, index: int) -> ColumnDimension:
    """Resolve a column dimension even when OOXML stores a range.

    Returns:
        The matching column-dimension object.

    """
    matching = next(
        (
            dimension
            for dimension in worksheet.column_dimensions.values()
            if dimension.min <= index <= dimension.max
        ),
        None,
    )
    if matching is not None:
        return matching
    return worksheet.column_dimensions[get_column_letter(index)]


def grouped_column_failures(
    worksheet: Worksheet,
    visible_count: int,
    total_count: int,
    label: str,
) -> list[str]:
    """Check one visible input surface followed by one collapsed detail group.

    Returns:
        Every grouped-column contract violation.

    """
    failures = [
        f"{label} core column {index} is unexpectedly hidden"
        for index in range(1, visible_count + 1)
        if column_dimension_for(worksheet, index).hidden
    ]
    failures.extend(
        f"{label} detail column {index} is not hidden at outline level 1"
        for index in range(visible_count + 1, total_count + 1)
        if not column_dimension_for(worksheet, index).hidden
        or column_dimension_for(worksheet, index).outlineLevel != 1
    )
    summary = column_dimension_for(worksheet, total_count + 1)
    if not summary.collapsed:
        failures.append(f"{label} hidden group has no collapsed [+] summary column")
    return failures
