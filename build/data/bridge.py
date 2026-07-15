"""Provider-neutral describe and planning orchestration for workbook agents."""

from __future__ import annotations

import hashlib
from collections.abc import Mapping
from copy import deepcopy
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import TYPE_CHECKING
from zipfile import BadZipFile
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from openpyxl.utils.exceptions import InvalidFileException

from ..spec.capacity import CONFIG_ROWS, DATA_ROWS
from ..spec.items import ITEMS_COLUMNS, RAID_COLUMNS
from .contract import (
    CHANGESET_SCHEMA,
    CONTRACT_NAME,
    CONTRACT_VERSION,
    ITEM_WRITABLE_FIELDS,
    RAID_WRITABLE_FIELDS,
    ContractError,
    canonical_json,
    parse_changeset,
)
from .diagnostics import Diagnostic, DiagnosticSeverity
from .export import ExportError, ExportResult, export_workbook
from .inject import InjectError, Reconciliation, validate_snapshot
from .merge import MergeResult, merge_changes
from .schema import (
    DATA_TABLES,
    SETTINGS_DESCRIPTIONS,
    SETTINGS_TYPES,
    TABLES_BY_NAME,
    schema_fingerprint,
)

if TYPE_CHECKING:
    from collections.abc import Sequence

    from .schema import TableSchema
    from .snapshot import Snapshot

_OWNERSHIP = {
    "I": "input",
    "S": "source_identity",
    "V": "vba",
    "F": "formula",
}


@dataclass(frozen=True, kw_only=True, slots=True)
class PlanEvaluation:
    """One complete planning evaluation for public output or approved apply."""

    result: dict[str, object]
    change_set: dict[str, object] | None
    base_snapshot: Snapshot | None
    intended_snapshot: Snapshot | None
    target: dict[str, str] | None
    token: str | None
    valid: bool
    conflict: bool
    has_changes: bool


@dataclass(frozen=True, kw_only=True, slots=True)
class _EffectiveTime:
    date: date
    timezone: str | None
    utc_offset: str
    expires_at: str

    def as_dict(self) -> dict[str, str | None]:
        """Return the complete approval-time boundary.

        Returns:
            The date, zone, offset and expiry fields bound into approval.

        """
        return {
            "date": self.date.isoformat(),
            "timezone": self.timezone,
            "utc_offset": self.utc_offset,
            "expires_at": self.expires_at,
        }


def _json_value(value: object) -> object:
    if isinstance(value, date):
        return value.isoformat()
    return value


def _serialize_rows(
    table_schema: TableSchema,
    rows: Sequence[Mapping[str, object]],
) -> list[dict[str, object]]:
    known = table_schema.column_names
    documents: list[dict[str, object]] = []
    for row in rows:
        document = {name: _json_value(row.get(name)) for name in known}
        for name in sorted(set(row) - set(known)):
            document[name] = _json_value(row[name])
        documents.append(document)
    return documents


def _value_type(spec: Mapping[str, object]) -> str:
    fmt = spec.get("fmt")
    if fmt == "date":
        return "date"
    if fmt == "int":
        return "integer"
    return "text"


def _column_contract(
    specs: Sequence[Mapping[str, object]],
    writable: frozenset[str],
) -> list[dict[str, object]]:
    return [
        {
            "name": str(spec["name"]),
            "value_type": _value_type(spec),
            "ownership": _OWNERSHIP[str(spec["kind"])],
            "writable": str(spec["name"]) in writable,
        }
        for spec in specs
    ]


def _clock_instant(now: datetime | None) -> datetime:
    if now is None:
        return datetime.now().astimezone()
    return now if now.tzinfo is not None else now.astimezone()


def _system_zone_name() -> str | None:
    try:
        resolved = str(Path("/etc/localtime").resolve(strict=True))
    except OSError:
        return None
    marker = "/zoneinfo/"
    if marker not in resolved:
        return None
    candidate = resolved.split(marker, maxsplit=1)[1]
    try:
        ZoneInfo(candidate)
    except ZoneInfoNotFoundError:
        return None
    return candidate


def _local_zone_name(instant: datetime, *, inspect_system: bool) -> str | None:
    key = getattr(instant.tzinfo, "key", None)
    if isinstance(key, str) and key:
        return key
    if instant.tzname() in {"UTC", "GMT"} and instant.utcoffset() == timedelta(0):
        return "UTC"
    return _system_zone_name() if inspect_system else None


def _offset_text(instant: datetime) -> str:
    offset = instant.utcoffset()
    if offset is None:
        message = "effective planning time must have a UTC offset"
        raise ValueError(message)
    total_minutes = int(offset.total_seconds() // 60)
    sign = "+" if total_minutes >= 0 else "-"
    hours, minutes = divmod(abs(total_minutes), 60)
    return f"{sign}{hours:02d}:{minutes:02d}"


def _effective_time(now: datetime | None) -> _EffectiveTime:
    instant = _clock_instant(now)
    zone_name = _local_zone_name(instant, inspect_system=now is None)
    midnight_zone = instant.tzinfo
    if zone_name is not None:
        try:
            midnight_zone = ZoneInfo(zone_name)
        except ZoneInfoNotFoundError:
            zone_name = None
    tomorrow = instant.date() + timedelta(days=1)
    expires = datetime.combine(tomorrow, time.min, tzinfo=midnight_zone)
    return _EffectiveTime(
        date=instant.date(),
        timezone=zone_name,
        utc_offset=_offset_text(instant),
        expires_at=expires.isoformat(timespec="seconds"),
    )


def _effective_date(now: datetime | None) -> str:
    return _effective_time(now).date.isoformat()


def _target(exported: ExportResult) -> dict[str, str]:
    return {
        "workbook_sha256": exported.snapshot.workbook_digest,
        "workbook_schema_fingerprint": exported.workbook_schema_fingerprint,
        "build_schema_fingerprint": schema_fingerprint(),
    }


def _diagnostic(
    code: str,
    severity: DiagnosticSeverity,
    message: str,
    hint: str,
    *,
    pointer: str = "",
) -> Diagnostic:
    return Diagnostic(
        code=code,
        severity=severity,
        phase="plan",
        pointer=pointer,
        operation_id=None,
        message=message,
        hint=hint,
    )


def _target_diagnostics(
    requested: Mapping[str, object],
    actual: Mapping[str, str],
) -> tuple[Diagnostic, ...]:
    labels = {
        "workbook_sha256": "workbook digest",
        "workbook_schema_fingerprint": "observed workbook schema fingerprint",
        "build_schema_fingerprint": "build-source schema fingerprint",
    }
    diagnostics: list[Diagnostic] = []
    for target_field, label in labels.items():
        if requested.get(target_field) == actual[target_field]:
            continue
        diagnostics.append(
            _diagnostic(
                f"target.{target_field}",
                "error",
                f"The change set's {label} does not match the planning target.",
                "Run describe again and rebuild the change set against its exact target values.",
                pointer=f"/target/{target_field}",
            )
        )
    return tuple(diagnostics)


def _migration_adjustments(
    exported: ExportResult,
    reconciliation: Reconciliation,
) -> tuple[str, ...]:
    adjustments = [
        f"{table}: add current column {column!r} during rebuild"
        for table, columns in sorted(exported.added_columns.items())
        for column in columns
    ]
    adjustments.extend(
        f"{name}: apply current default {reconciliation.settings[name]!r}"
        for name in reconciliation.defaulted_settings
    )
    adjustments.extend(
        f"{name}: bump {before} to {after} above existing identifiers"
        for name, (before, after) in sorted(reconciliation.counter_bumps.items())
    )
    adjustments.extend(reconciliation.adjustments)
    return tuple(adjustments)


def _planning_warnings(
    exported: ExportResult,
    reconciliation: Reconciliation,
) -> tuple[Diagnostic, ...]:
    warnings = [
        _diagnostic(
            "workbook.reconciliation_warning",
            "warning",
            message,
            "Review the existing workbook value; it is preserved unless an operation changes it.",
        )
        for message in reconciliation.warnings
    ]
    warnings.extend(
        _diagnostic(
            "workbook.export_note",
            "warning",
            message,
            "Review the workbook note before approving this plan.",
        )
        for message in exported.notes
    )
    warnings.extend(
        _diagnostic(
            "workbook.examples_skipped",
            "warning",
            f"{table}: skipped {count} marked example row(s) while reading authored data.",
            "Remove example markers from any row that should be treated as authored data.",
        )
        for table, count in sorted(exported.skipped_examples.items())
    )
    return tuple(warnings)


def _snapshot_state(snapshot: Snapshot) -> dict[str, object]:
    return {
        "schema_fingerprint": snapshot.schema_fingerprint,
        "settings": dict(snapshot.settings),
        "tables": {table: [dict(row) for row in rows] for table, rows in snapshot.tables.items()},
    }


def _summary(merge: MergeResult | None) -> dict[str, int]:
    actions = {"create": 0, "update": 0, "mark_deleted": 0, "noop": 0}
    changes = 0
    if merge is not None:
        for operation in merge.operations:
            actions[operation.action] += 1
            changes += len(operation.diffs)
    return {**actions, "field_changes": changes}


@dataclass(frozen=True, kw_only=True, slots=True)
class _PlanPresentation:
    effective: _EffectiveTime
    target: Mapping[str, str] | None = None
    merge: MergeResult | None = None
    adjustments: Sequence[str] = ()
    warnings: Sequence[Diagnostic] = ()
    errors: Sequence[Diagnostic] = ()
    token: str | None = None
    conflict: bool = False

    @property
    def has_changes(self) -> bool:
        """Return whether rebuild would alter authored or workbook structure."""
        if self.adjustments:
            return True
        if self.merge is None:
            return False
        return any(operation.action != "noop" for operation in self.merge.operations)

    def as_dict(self) -> dict[str, object]:
        """Return the stable public plan document.

        Returns:
            Exact operation diffs, diagnostics, time binding and optional token.

        """
        operations = (
            []
            if self.merge is None
            else [operation.as_dict() for operation in self.merge.operations]
        )
        document: dict[str, object] = {
            "result": "plan",
            "valid": not self.errors,
            "conflict": self.conflict,
            "target": None if self.target is None else dict(self.target),
            "effective_time": self.effective.as_dict(),
            "summary": _summary(self.merge),
            "operations": operations,
            "migration_adjustments": list(self.adjustments),
            "warnings": [diagnostic.as_dict() for diagnostic in self.warnings],
            "errors": [diagnostic.as_dict() for diagnostic in self.errors],
            "has_changes": self.has_changes,
        }
        if self.token is not None:
            document["plan_token"] = self.token
        return document


def describe_workbook(
    workbook: str | Path,
    *,
    now: datetime | None = None,
) -> dict[str, object]:
    """Return the deterministic agent contract and current authored state.

    Returns:
        A JSON-compatible description containing the mutation schema, target
        digests, Config choices, ownership metadata and ordered current rows.

    """
    source = Path(workbook).expanduser().resolve()
    exported = export_workbook(source)
    snapshot = exported.snapshot
    config_tables = {
        table.table: _serialize_rows(
            table,
            snapshot.tables.get(table.table, ()),
        )
        for table in DATA_TABLES
        if table.sheet == "Config"
    }
    settings = [
        {
            "name": name,
            "value": _json_value(value),
            "value_type": SETTINGS_TYPES.get(name, "unknown"),
            "description": SETTINGS_DESCRIPTIONS.get(name, ""),
        }
        for name, value in snapshot.settings.items()
    ]
    return {
        "result": "describe",
        "contract": {
            "name": CONTRACT_NAME,
            "version": CONTRACT_VERSION,
            "schema": deepcopy(CHANGESET_SCHEMA),
        },
        "target": _target(exported),
        "effective_date": _effective_date(now),
        "capacities": {
            "items": DATA_ROWS,
            "raid": DATA_ROWS,
            "config": CONFIG_ROWS,
        },
        "writable_fields": {
            "item": list(ITEM_WRITABLE_FIELDS),
            "raid": list(RAID_WRITABLE_FIELDS),
        },
        "columns": {
            "item": _column_contract(ITEMS_COLUMNS, frozenset(ITEM_WRITABLE_FIELDS)),
            "raid": _column_contract(RAID_COLUMNS, frozenset(RAID_WRITABLE_FIELDS)),
        },
        "config": {
            "settings": settings,
            "tables": config_tables,
        },
        "records": {
            "items": _serialize_rows(
                TABLES_BY_NAME["tblItems"],
                snapshot.tables.get("tblItems", ()),
            ),
            "raid": _serialize_rows(
                TABLES_BY_NAME["tblRAID"],
                snapshot.tables.get("tblRAID", ()),
            ),
        },
        "schema_notes": {
            "added_columns": {
                table: list(columns) for table, columns in sorted(exported.added_columns.items())
            },
            "unknown_columns": {
                table: list(columns) for table, columns in sorted(exported.unknown_columns.items())
            },
            "skipped_examples": dict(sorted(exported.skipped_examples.items())),
            "notes": list(exported.notes),
        },
    }


@dataclass(slots=True)
class _Planner:
    payload: bytes
    source: Path
    effective: _EffectiveTime
    change_set: dict[str, object] | None = field(default=None, init=False)
    exported: ExportResult | None = field(default=None, init=False)
    target: dict[str, str] | None = field(default=None, init=False)
    reconciliation: Reconciliation | None = field(default=None, init=False)

    def _failure(
        self,
        errors: Sequence[Diagnostic],
        *,
        conflict: bool = False,
    ) -> PlanEvaluation:
        presentation = _PlanPresentation(
            effective=self.effective,
            target=self.target,
            errors=errors,
            conflict=conflict,
        )
        base = None if self.exported is None else self.exported.snapshot
        return PlanEvaluation(
            result=presentation.as_dict(),
            change_set=self.change_set,
            base_snapshot=base,
            intended_snapshot=None,
            target=self.target,
            token=None,
            valid=False,
            conflict=conflict,
            has_changes=False,
        )

    def _parse(self) -> PlanEvaluation | None:
        try:
            self.change_set = parse_changeset(self.payload)
        except ContractError as error:
            return self._failure(error.diagnostics)
        return None

    def _read(self) -> PlanEvaluation | None:
        try:
            self.exported = export_workbook(self.source)
        except (BadZipFile, EOFError, ExportError, InvalidFileException, OSError) as error:
            diagnostic = _diagnostic(
                "workbook.unreadable",
                "error",
                f"The target workbook could not be read: {error}",
                "Supply an existing, structurally valid Excel project-management workbook.",
            )
            return self._failure((diagnostic,))

        self.target = _target(self.exported)
        requested_target = {} if self.change_set is None else self.change_set.get("target")
        requested = requested_target if isinstance(requested_target, Mapping) else {}
        target_errors = _target_diagnostics(requested, self.target)
        if target_errors:
            return self._failure(target_errors, conflict=True)
        return self._recheck_digest()

    def _recheck_digest(self) -> PlanEvaluation | None:
        if self.exported is None:
            message = "planner digest recheck requires an exported workbook"
            raise RuntimeError(message)
        try:
            digest_after_read = hashlib.sha256(self.source.read_bytes()).hexdigest()
        except OSError as error:
            diagnostic = _diagnostic(
                "target.workbook_recheck",
                "error",
                f"The workbook could not be rechecked after planning read: {error}",
                "Restore access to the workbook and run plan again.",
            )
            return self._failure((diagnostic,), conflict=True)
        if digest_after_read == self.exported.snapshot.workbook_digest:
            return None
        diagnostic = _diagnostic(
            "target.workbook_changed_during_plan",
            "error",
            "The workbook changed while its authored state was being read.",
            "Close competing edits, run describe again and re-create the change set.",
        )
        return self._failure((diagnostic,), conflict=True)

    def _reconcile(self) -> PlanEvaluation | None:
        if self.exported is None:
            message = "planner reconciliation requires an exported workbook"
            raise RuntimeError(message)
        try:
            self.reconciliation = validate_snapshot(self.exported.snapshot)
        except InjectError as error:
            diagnostic = _diagnostic(
                "workbook.unmigratable",
                "error",
                f"The workbook cannot be migrated without data loss: {error}",
                "Correct the reported workbook structure or authored data before planning again.",
            )
            return self._failure((diagnostic,))
        return None

    def _token(
        self,
        intended_snapshot: Snapshot,
        warnings: Sequence[Diagnostic],
        adjustments: Sequence[str],
    ) -> str:
        if self.change_set is None or self.target is None:
            message = "planner token requires a parsed change set and target"
            raise RuntimeError(message)
        binding = {
            "format": 1,
            "change_set": self.change_set,
            "target": self.target,
            "intended_authored_state": _snapshot_state(intended_snapshot),
            "warnings": [warning.as_dict() for warning in warnings],
            "migration_adjustments": list(adjustments),
            "effective_time": self.effective.as_dict(),
        }
        return hashlib.sha256(canonical_json(binding)).hexdigest()

    def _finish(self) -> PlanEvaluation:
        if (
            self.change_set is None
            or self.exported is None
            or self.target is None
            or self.reconciliation is None
        ):
            message = "planner finish requires parsed, exported and reconciled state"
            raise RuntimeError(message)
        adjustments = _migration_adjustments(self.exported, self.reconciliation)
        merge = merge_changes(
            self.reconciliation.normalized_snapshot,
            self.change_set,
            effective_date=self.effective.date,
        )
        warnings = (
            *_planning_warnings(self.exported, self.reconciliation),
            *(item for item in merge.diagnostics if item.severity == "warning"),
        )
        errors = tuple(item for item in merge.diagnostics if item.severity == "error")
        valid = not errors
        token = self._token(merge.snapshot, warnings, adjustments) if valid else None
        presentation = _PlanPresentation(
            effective=self.effective,
            target=self.target,
            merge=merge,
            adjustments=adjustments,
            warnings=warnings,
            errors=errors,
            token=token,
        )
        result = presentation.as_dict()
        return PlanEvaluation(
            result=result,
            change_set=self.change_set,
            base_snapshot=self.exported.snapshot,
            intended_snapshot=merge.snapshot if valid else None,
            target=self.target,
            token=token,
            valid=valid,
            conflict=False,
            has_changes=presentation.has_changes,
        )

    def evaluate(self) -> PlanEvaluation:
        """Run strict parsing, target checks, reconciliation and merge.

        Returns:
            A complete public and internal plan evaluation.

        """
        for stage in (self._parse, self._read, self._reconcile):
            failure = stage()
            if failure is not None:
                return failure
        return self._finish()


def evaluate_plan(
    payload: bytes,
    workbook: str | Path,
    *,
    now: datetime | None = None,
) -> PlanEvaluation:
    """Parse and evaluate a read-only, approval-bound workbook plan.

    Returns:
        The public plan document plus the typed state needed by approved apply.

    """
    planner = _Planner(
        payload=payload,
        source=Path(workbook).expanduser().resolve(),
        effective=_effective_time(now),
    )
    return planner.evaluate()


def plan_workbook(
    payload: bytes,
    workbook: str | Path,
    *,
    now: datetime | None = None,
) -> dict[str, object]:
    """Return one strict, read-only workbook change plan.

    Returns:
        A JSON-compatible plan with a token only when every check succeeds.

    """
    return evaluate_plan(payload, workbook, now=now).result


__all__ = ["PlanEvaluation", "describe_workbook", "evaluate_plan", "plan_workbook"]
